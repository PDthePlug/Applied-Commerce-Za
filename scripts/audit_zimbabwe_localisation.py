from __future__ import annotations

from collections import Counter
from pathlib import Path
from docx import Document
import argparse
import json
import re

PATTERNS = {
    "south_africa": re.compile(r"\bSouth Africa\b|\bSouth African\b", re.I),
    "dbe": re.compile(r"\bDBE\b|Department of Basic Education", re.I),
    "currency": re.compile(r"(?<![A-Za-z])R\s?\d|\brand(?:s)?\b|\bZAR\b", re.I),
    "stokvel": re.compile(r"\bstokvel(?:s)?\b", re.I),
    "spaza": re.compile(r"\bspaza\b", re.I),
    "taxi_rank": re.compile(r"\btaxi rank\b", re.I),
    "matric": re.compile(r"\bmatric\b", re.I),
    "nsfas": re.compile(r"\bNSFAS\b", re.I),
    "sars": re.compile(r"\bSARS\b", re.I),
}

PLACE_RE = re.compile(
    r"\b(?:Johannesburg|Gauteng|Durban|Soweto|Umlazi|Katlehong|Alexandra|Cape Town|Pretoria)\b",
    re.I,
)


def document_text(path: Path) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells if cell.text)
    return "\n".join(chunks)


def audit(path: Path) -> dict[str, int]:
    text = document_text(path)
    counts = {name: len(regex.findall(text)) for name, regex in PATTERNS.items()}
    counts["sa_places"] = len(PLACE_RE.findall(text))
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description="Inventory South Africa-specific context in Applied Commerce manuscripts.")
    parser.add_argument("source_dir", type=Path, help="Directory containing APPLIED COMMERCE Grade *.docx")
    parser.add_argument("--json", action="store_true", help="Print machine-readable JSON")
    args = parser.parse_args()

    rows = {}
    for path in sorted(args.source_dir.glob("APPLIED COMMERCE Grade *.docx")):
        match = re.search(r"Grade (\d+)", path.name)
        if not match:
            continue
        rows[int(match.group(1))] = audit(path)

    if args.json:
        print(json.dumps(rows, indent=2))
        return

    if not rows:
        raise SystemExit("No Applied Commerce source manuscripts found.")

    headers = ["grade", *PATTERNS.keys(), "sa_places"]
    print("\t".join(headers))
    for grade, counts in rows.items():
        print("\t".join([str(grade), *[str(counts[h]) for h in headers[1:]]]))


if __name__ == "__main__":
    main()
